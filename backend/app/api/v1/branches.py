import io
import json

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import require_roles_with_matrix
from app.db.session import get_db
from app.models.company import CompanyMaster
from app.services.company_service import CompanyService

router = APIRouter(prefix="/admin-config/branches", tags=["Branch Master"])
view_access = require_roles_with_matrix([1, 2], ["Masters", "Masters (State/City/Scope)"])
edit_access = require_roles_with_matrix([1], ["Masters", "Masters (State/City/Scope)"])
COLUMNS = [
    "Company Name",
    "Branch Name",
    "Branch ID",
    "City",
    "State",
    "Country",
    "Address",
    "IC Member 1 Employee ID",
    "IC Member 2 Employee ID",
    "IC Member 3 Employee ID",
    "IC Member 4 Employee ID",
]


class BranchInput(BaseModel):
    company_id: int
    branch_name: str = Field(min_length=1, max_length=150)
    branch_id: str = Field(min_length=1, max_length=50)
    city: str = Field(default="", max_length=150)
    state: str = Field(default="", max_length=150)
    country: str = Field(default="", max_length=150)
    address1: str = Field(default="", max_length=1000)
    ic_member_ids: list[str] = Field(default_factory=list, max_length=4)
    original_branch_id: str | None = None


def branch_rows(company):
    try:
        rows = json.loads(company.branches_json or "[]")
    except (TypeError, ValueError):
        raise HTTPException(409, "Existing branch data needs correction before editing.")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise HTTPException(409, "Invalid existing branch data.")
    return rows


async def employees_for(db, company_id):
    result = await db.execute(
        text(
            """
        SELECT employee_id, first_name, last_name, ic_role
        FROM posh_employee_master WHERE company_id = :company_id AND status = 'Active'
        UNION ALL
        SELECT employee_id, first_name, last_name, ic_role
        FROM user_master WHERE company_id = :company_id AND status = 'Active' AND is_deleted = 'N'
    """
        ),
        {"company_id": company_id},
    )
    employees = {}
    for row in result.mappings():
        employees[row["employee_id"]] = {
            "employee_id": row["employee_id"],
            "name": " ".join(filter(None, [row["first_name"], row["last_name"]])),
            "ic_role": row["ic_role"] or employees.get(row["employee_id"], {}).get("ic_role", ""),
        }
    return list(employees.values())


def merge_branch(rows, data, employees):
    payload = data.model_dump(exclude={"company_id", "original_branch_id"})
    payload = {
        key: value.strip() if isinstance(value, str) else value for key, value in payload.items()
    }
    if not payload["branch_name"] or not payload["branch_id"]:
        raise HTTPException(422, "Branch name and ID are required.")
    ids = [value.strip() for value in data.ic_member_ids if value.strip()]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Select each IC member only once.")
    eligible = {employee["employee_id"] for employee in employees}
    if any(value not in eligible for value in ids):
        raise HTTPException(422, "IC members must be active employees of the selected company.")
    payload["ic_member_ids"] = ids
    original = data.original_branch_id
    index = (
        next((i for i, row in enumerate(rows) if row.get("branch_id") == original), None)
        if original is not None
        else None
    )
    if original is not None and index is None:
        raise HTTPException(404, "Branch no longer exists. Refresh and try again.")
    if original is not None and payload["branch_id"] != original:
        raise HTTPException(
            422, "Existing branch IDs cannot be changed because reports reference them."
        )
    if any(
        i != index
        and str(row.get("branch_id", "")).strip().casefold() == payload["branch_id"].casefold()
        for i, row in enumerate(rows)
    ):
        raise HTTPException(409, "Branch ID already exists for this company.")
    updated = [dict(row) for row in rows]
    if index is None:
        updated.append(payload)
    else:
        updated[index] = {**updated[index], **payload}
    return updated


@router.get("")
async def list_branches(db: AsyncSession = Depends(get_db), current_user=Depends(view_access)):
    companies = await CompanyService().get_visible_for_user(db, current_user)
    output = []
    for company in companies:
        employees = await employees_for(db, company.company_id)
        output.append(
            {
                "company_id": company.company_id,
                "company_name": company.company_name,
                "approval_status": company.approval_status,
                "branches": branch_rows(company),
                "employees": employees,
                "presiding_officers": [
                    person for person in employees if "presiding" in person["ic_role"].lower()
                ],
            }
        )
    return output


@router.post("")
async def save_branch(
    data: BranchInput, db: AsyncSession = Depends(get_db), current_user=Depends(edit_access)
):
    result = await db.execute(
        select(CompanyMaster)
        .where(CompanyMaster.company_id == data.company_id, CompanyMaster.is_deleted == "N")
        .with_for_update()
    )
    company = result.scalar_one_or_none()
    if not company:
        raise HTTPException(404, "Company not found.")
    company.branches_json = json.dumps(
        merge_branch(branch_rows(company), data, await employees_for(db, company.company_id))
    )
    await db.commit()
    return {"detail": "Branch saved."}


@router.get("/template")
async def branch_template(current_user=Depends(view_access)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Branches"
    sheet.append(COLUMNS)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
    stream = io.BytesIO()
    workbook.save(stream)
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="branch_master_template.xlsx"'},
    )


@router.post("/bulk-upload")
async def bulk_upload(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(edit_access),
):
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Workbook must be smaller than 5 MB.")
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheet = workbook.active
        records = sheet.iter_rows(values_only=True)
        if list(next(records, ())) != COLUMNS:
            raise ValueError("Use the Branch Master Excel template with unchanged column headings.")
        result = await db.execute(
            select(CompanyMaster)
            .where(CompanyMaster.is_deleted == "N", CompanyMaster.approval_status == "Approved")
            .order_by(CompanyMaster.company_id)
            .with_for_update()
        )
        companies = result.scalars().all()
        pending = {}
        people = {}
        count = 0
        for number, row in enumerate(records, 2):
            if number > 1001:
                raise ValueError("Maximum 1,000 rows per upload.")
            if not any(value is not None for value in row):
                continue
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(value.startswith("=") for value in values):
                raise ValueError(f"Row {number}: use values, not Excel formulas.")
            values += [""] * (len(COLUMNS) - len(values))
            matching = [
                company
                for company in companies
                if company.company_name.casefold() == values[0].casefold()
            ]
            if len(matching) != 1:
                raise ValueError(f"Row {number}: company name must match one approved company.")
            company = matching[0]
            if company.company_id not in people:
                people[company.company_id] = await employees_for(db, company.company_id)
            data = BranchInput(
                company_id=company.company_id,
                branch_name=values[1],
                branch_id=values[2],
                city=values[3],
                state=values[4],
                country=values[5],
                address1=values[6],
                ic_member_ids=[value for value in values[7:11] if value],
            )
            try:
                pending[company.company_id] = merge_branch(
                    pending.get(company.company_id, branch_rows(company)),
                    data,
                    people[company.company_id],
                )
            except HTTPException as exc:
                raise ValueError(f"Row {number}: {exc.detail}") from exc
            count += 1
        if not count:
            raise ValueError("Workbook contains no branches.")
        for company in companies:
            if company.company_id in pending:
                company.branches_json = json.dumps(pending[company.company_id])
        await db.commit()
        return {"detail": f"Created {count} branches."}
    except (ValueError, KeyError, TypeError) as exc:
        await db.rollback()
        raise HTTPException(422, str(exc)) from exc
    except HTTPException:
        await db.rollback()
        raise
    except Exception as exc:
        await db.rollback()
        raise HTTPException(
            422, "Unable to import workbook. Check the Excel template and data."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
