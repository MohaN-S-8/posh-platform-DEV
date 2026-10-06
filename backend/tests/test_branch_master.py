import io
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from app.api.v1 import branches
from app.api.v1.branches import BranchInput, merge_branch


def test_branch_rejects_member_from_another_company():
    data = BranchInput(company_id=2, branch_name="Main", branch_id="BR1", ic_member_ids=["OTHER"])
    with pytest.raises(HTTPException) as error:
        merge_branch([], data, [{"employee_id": "OWN"}])
    assert error.value.status_code == 422


def test_duplicate_branch_ids_are_case_insensitive():
    data = BranchInput(company_id=2, branch_name="Second", branch_id="br1")
    with pytest.raises(HTTPException) as error:
        merge_branch([{"branch_id": "BR1"}], data, [])
    assert error.value.status_code == 409


def test_edit_preserves_company_setup_address_and_unrelated_branches():
    rows = [
        {"branch_id": "BR1", "address2": "Floor 2"},
        {"branch_id": "BR2", "branch_name": "Other"},
    ]
    data = BranchInput(
        company_id=2, branch_name="Updated", branch_id="BR1", original_branch_id="BR1"
    )
    updated = merge_branch(rows, data, [])
    assert updated[0]["address2"] == "Floor 2"
    assert updated[0]["branch_name"] == "Updated"
    assert updated[1] == rows[1]
    assert "branch_name" not in rows[0]


def test_branch_id_cannot_change_after_creation():
    data = BranchInput(company_id=2, branch_name="Main", branch_id="NEW", original_branch_id="BR1")
    with pytest.raises(HTTPException) as error:
        merge_branch([{"branch_id": "BR1"}], data, [])
    assert error.value.status_code == 422


def test_duplicate_committee_member_is_rejected():
    data = BranchInput(
        company_id=2, branch_name="Main", branch_id="BR1", ic_member_ids=["OWN", "OWN"]
    )
    with pytest.raises(HTTPException):
        merge_branch([], data, [{"employee_id": "OWN"}])


def upload_workbook(rows):
    workbook = Workbook()
    workbook.active.append(branches.COLUMNS)
    for row in rows:
        workbook.active.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return UploadFile(filename="branches.xlsx", file=stream)


def database_with(company):
    db = AsyncMock()
    result = MagicMock()
    result.scalars.return_value.all.return_value = [company]
    db.execute.return_value = result
    return db


@pytest.mark.asyncio
async def test_template_is_valid_excel():
    response = await branches.branch_template(current_user=None)
    workbook = load_workbook(io.BytesIO(response.body))
    assert [cell.value for cell in workbook.active[1]] == branches.COLUMNS
    assert workbook.active.freeze_panes == "A2"
    workbook.close()


@pytest.mark.asyncio
async def test_bulk_upload_adds_rows_and_preserves_existing_branches(monkeypatch):
    company = SimpleNamespace(
        company_id=2, company_name="Alpha", branches_json='[{"branch_id":"OLD"}]'
    )
    db = database_with(company)
    monkeypatch.setattr(branches, "employees_for", AsyncMock(return_value=[{"employee_id": "OWN"}]))
    file = upload_workbook([["Alpha", "Main", "BR1", "Chennai", "TN", "IN", "Road 1", "OWN"]])
    response = await branches.bulk_upload(file=file, db=db, current_user=None)
    assert response["detail"] == "Created 1 branches."
    assert json.loads(company.branches_json)[0] == {"branch_id": "OLD"}
    assert json.loads(company.branches_json)[1]["ic_member_ids"] == ["OWN"]
    db.commit.assert_awaited_once()


@pytest.mark.asyncio
async def test_invalid_bulk_row_does_not_partially_save(monkeypatch):
    company = SimpleNamespace(company_id=2, company_name="Alpha", branches_json="[]")
    db = database_with(company)
    monkeypatch.setattr(branches, "employees_for", AsyncMock(return_value=[]))
    file = upload_workbook([["Alpha", "Main", "BR1"], ["Unknown Company", "Second", "BR2"]])
    with pytest.raises(HTTPException) as error:
        await branches.bulk_upload(file=file, db=db, current_user=None)
    assert "Row 3" in error.value.detail
    assert company.branches_json == "[]"
    db.commit.assert_not_awaited()
    db.rollback.assert_awaited_once()
